"""Render each docs/*.xlsx into a clean, gridline-free HTML page under site/.

Usage: python build_site.py
Re-run after editing any workbook. Values are the cached results Excel/Sheets saved.
"""
import datetime as dt
import html
import re
import shutil
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

ROOT = Path(__file__).parent
DOCS, SITE = ROOT / "docs", ROOT / "site"
OWNER = "Jeuz Salvedia"
# Clicking a cell shows its formula in the formula bar, like Excel. Set False to show only displayed values.
SHOW_FORMULAS = True
# Order on the hub page; first sheet of every file is the output sheet.
FILES = [
    "Comparable Company Analysis",
    "Precedent Transaction Analysis",
    "Discounted Cash Flow Analysis",
    "Merger Consequences Analysis",
]

# ---------- number formats ----------

DATE_TOKENS = [("yyyy", "%Y"), ("yy", "%y"), ("mmmm", "%B"), ("mmm", "%b"), ("mm", "%m"), ("m", "%-m"), ("dd", "%d"), ("d", "%-d")]


def fmt_date(v, fmt):
    out, i = "", 0
    while i < len(fmt):
        for tok, code in DATE_TOKENS:
            if fmt.startswith(tok, i):
                s = v.strftime(code.replace("-", ""))
                out += s.lstrip("0") if "-" in code else s
                i += len(tok)
                break
        else:
            out += fmt[i]
            i += 1
    return out


def fmt_general(v):
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return str(int(v))
    if isinstance(v, float):
        return f"{v:.10g}" if abs(v) >= 1e-4 or v == 0 else f"{v:.4E}"
    return str(v)


def fmt_number(v, fmt):
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (dt.datetime, dt.date)):
        return fmt_date(v, fmt) if fmt != "General" else v.strftime("%Y-%m-%d")
    if not isinstance(v, (int, float)):
        return str(v)
    if fmt == "General":
        return fmt_general(v)
    sections = fmt.split(";")
    if v < 0 and len(sections) > 1:
        sec, v = sections[1], -v  # negative section carries its own sign/parens
    elif v == 0 and len(sections) > 2:
        sec = sections[2]
    else:
        sec = sections[0]
    sec = re.sub(r"\[\$([^\]-]*)[^\]]*\]", r"\1", sec)  # [$₱] -> ₱
    sec = re.sub(r"\[[^\]]*\]", "", sec)  # drop colors/conditions
    m = re.search(r"[#0][#0,]*(\.0+)?", sec)
    if not m:
        return sec.replace('"', "")
    decimals = len(m.group(1)) - 1 if m.group(1) else 0
    pct = "%" in sec
    num = f"{abs(v) * (100 if pct else 1):{',' if ',' in m.group(0) else ''}.{decimals}f}"
    sign = "-" if v < 0 and len(sections) == 1 else ""  # Excel puts the sign before the currency symbol
    pre, post = sec[: m.start()], sec[m.end():]
    return (sign + pre + num + post).replace('"', "").replace("\\", "")


def selfcheck():
    from openpyxl.styles import Border, PatternFill, Side
    box = openpyxl.Workbook().active  # a box merged J4:N4: text and fill in J4, only borders beyond it
    box["J4"] = "Target Description"
    box["J4"].fill = PatternFill("solid", fgColor="CC0000")
    box.merge_cells("J4:N4")
    box["N4"].border = Border(right=Side("thin"))
    box["N6"].border = Border(bottom=Side("thin"))
    grid = render_sheet(box, [], box)
    assert '<th data-c="N">' in grid and grid.count('<th class="rh">') == 6, "trimmed away a merged or bordered cell"
    from openpyxl.cell.text import InlineFont
    mixed = rich_text_html(CellRichText(TextBlock(InlineFont(b=True), "Coverage Ratios "), TextBlock(InlineFont(i=True), "(2)")), [])
    assert 'font-weight:700;font-style:normal">Coverage Ratios </span>' in mixed, mixed
    assert 'font-weight:400;font-style:italic">(2)</span>' in mixed, mixed
    assert fmt_number(7.4, "0.0x") == "7.4x"
    assert fmt_number(0.181, "0.0%") == "18.1%"
    assert fmt_number(1234.5, "[$₱]#,##0.00") == "₱1,234.50"
    assert fmt_number(-1234.5, "#,##0.00;(#,##0.00)") == "(1,234.50)"
    assert fmt_number(-2.5, "0.00") == "-2.50"
    assert fmt_number(-1614.7, "[$₱]#,##0.00") == "-₱1,614.70"
    assert fmt_number(0, '₱#,##0.00;(₱#,##0.00);"₱-"0.00') == "₱-0.00"
    assert fmt_number(2024.0, "General") == "2024"
    assert fmt_number(dt.datetime(2021, 11, 25), "mmmm d, yyyy") == "November 25, 2021"
    assert fmt_number(dt.datetime(2021, 3, 4), "dd/mm/yyyy") == "04/03/2021"


# ---------- colors ----------


def theme_colors(wb):
    if not wb.loaded_theme:
        return []
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme = ET.fromstring(wb.loaded_theme).find(".//a:clrScheme", ns)
    cols = []
    for el in scheme:
        c = el[0]
        cols.append(c.get("lastClr") or c.get("val") or "000000")
    cols[0], cols[1], cols[2], cols[3] = cols[1], cols[0], cols[3], cols[2]  # Excel index order: lt1, dk1, lt2, dk2
    return cols


def css_color(color, theme):
    if color is None:
        return None
    if color.type == "rgb" and isinstance(color.rgb, str):
        rgb = color.rgb[-6:]
        if color.rgb[:2] == "00" and len(color.rgb) == 8 and rgb == "000000":
            return None  # openpyxl's "unset" default
    elif color.type == "theme" and theme and color.theme < len(theme):
        rgb = theme[color.theme]
    else:
        return None
    r, g, b = (int(rgb[i:i + 2], 16) for i in (0, 2, 4))
    t = color.tint or 0
    if t:
        r, g, b = (round(c * (1 + t)) if t < 0 else round(c + (255 - c) * t) for c in (r, g, b))
    return f"#{r:02x}{g:02x}{b:02x}"


BORDER = {"thin": "1px solid", "hair": "1px solid", "medium": "2px solid", "thick": "3px solid", "double": "3px double",
          "dashed": "1px dashed", "mediumDashed": "2px dashed", "dotted": "1px dotted"}


def side_css(side, theme):
    if side is None or not side.style:
        return None
    return f"{BORDER.get(side.style, '1px solid')} {css_color(side.color, theme) or '#000'}"


# ---------- sheet rendering ----------


def rich_text_html(v, theme):
    """A cell whose text mixes fonts (e.g. a bold label with an italic "(2)" footnote): one span per run."""
    out = []
    for run in v:
        if not isinstance(run, TextBlock):  # plain run: inherits the cell's font
            out.append(html.escape(str(run)))
            continue
        f, st = run.font, []
        st.append(f"font-weight:{700 if f.b else 400}")
        st.append(f"font-style:{'italic' if f.i else 'normal'}")
        if f.u:
            st.append("text-decoration:underline")
        if f.sz:
            st.append(f"font-size:{f.sz:g}pt")
        fc = css_color(f.color, theme)
        if fc:
            st.append(f"color:{fc}")
        out.append(f'<span style="{";".join(st)}">{html.escape(run.text)}</span>')
    return "".join(out).replace("\n", "<br>")


def raw_value(v):
    """What Excel's formula bar shows for a cell: the formula, or the unformatted constant."""
    if isinstance(v, ArrayFormula):
        return "{" + (v.text or "") + "}"
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%m/%d/%Y")
    return fmt_general(v) if isinstance(v, float) else str(v)


def render_sheet(ws, theme, fws):
    # Trim trailing rows/cols that draw nothing (Sheets exports pad the used range). A cell draws something if
    # it has a value, a fill or a border; a merged range draws out to its far corner even when only its
    # top-left cell holds the text.
    def draws(c):
        b = c.border
        return (c.value is not None or (c.fill is not None and c.fill.fill_type == "solid")
                or (b is not None and any(getattr(getattr(b, s, None), "style", None) for s in ("left", "right", "top", "bottom"))))

    used = [(c.row, c.column) for row in ws.iter_rows() for c in row if draws(c)]
    used += [(rng.max_row, rng.max_col) for rng in ws.merged_cells.ranges]
    max_r, max_c = max(r for r, _ in used), max(c for _, c in used)
    merged, covered = {}, set()
    for rng in ws.merged_cells.ranges:
        merged[(rng.min_row, rng.min_col)] = rng
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    covered.add((r, c))

    widths = {}
    for dim in ws.column_dimensions.values():
        if dim.width and dim.min:
            for c in range(dim.min, (dim.max or dim.min) + 1):
                widths[c] = dim.width
    default_w = ws.sheet_format.defaultColWidth or 8.43
    total = 40
    cols = ['<col style="width:40px">']
    for c in range(1, max_c + 1):
        px = round(widths.get(c, default_w) * 100 / 12.63)  # Google Sheets: width 12.63 == 100px
        total += px
        cols.append(f'<col style="width:{px}px">')

    rows = []
    for r in range(1, max_r + 1):
        ht = ws.row_dimensions[r].ht
        tds = []
        for c in range(1, max_c + 1):
            if (r, c) in covered:
                continue
            cell = ws.cell(r, c)
            rng = merged.get((r, c))
            attrs, st = f' data-a="{get_column_letter(c)}{r}"', []
            fv = fws.cell(r, c).value
            if SHOW_FORMULAS and fv is not None:
                attrs += f' data-f="{html.escape(raw_value(fv))}"'
            right_cell, bottom_cell = cell, cell
            if rng:
                if rng.max_col > rng.min_col:
                    attrs += f' colspan="{rng.max_col - rng.min_col + 1}"'
                if rng.max_row > rng.min_row:
                    attrs += f' rowspan="{rng.max_row - rng.min_row + 1}"'
                right_cell = ws.cell(r, rng.max_col)
                bottom_cell = ws.cell(rng.max_row, c)
            if cell.fill is not None and cell.fill.fill_type == "solid":
                bg = css_color(cell.fill.fgColor, theme)
                if bg:
                    st.append(f"background:{bg}")
            for name, src in (("top", cell), ("left", cell), ("bottom", bottom_cell), ("right", right_cell)):
                b = src.border
                css = side_css(getattr(b, name, None), theme) if b is not None else None
                if css:
                    st.append(f"border-{name}:{css}")
            f = cell.font
            if f is not None:
                if f.b:
                    st.append("font-weight:700")
                if f.i:
                    st.append("font-style:italic")
                if f.u:
                    st.append("text-decoration:underline")
                if f.sz and f.sz != 10:
                    st.append(f"font-size:{f.sz:g}pt")
                fc = css_color(f.color, theme)
                if fc and fc != "#000000":
                    st.append(f"color:{fc}")
            v = cell.value
            is_num = isinstance(v, (int, float)) and not isinstance(v, bool)
            al = cell.alignment
            h = (al.horizontal if al is not None else None) or ("right" if is_num else None)
            if h in ("center", "centerContinuous"):
                st.append("text-align:center")
            elif h == "right":
                st.append("text-align:right")
            va = al.vertical if al is not None else None
            if va in ("top", "center"):
                st.append("vertical-align:" + ("middle" if va == "center" else "top"))
            if al is not None and al.wrap_text:
                st.append("white-space:pre-wrap")
            # Excel lets text spill into an empty neighbour; clip only when the neighbour has content.
            nxt = ws.cell(r, (rng.max_col if rng else c) + 1).value if c < max_c else 1
            if nxt is not None and not (al is not None and al.wrap_text):
                st.append("overflow:hidden")
            if isinstance(v, CellRichText):
                text = rich_text_html(v, theme)
            else:
                text = "" if v is None else html.escape(fmt_number(v, cell.number_format)).replace("\n", "<br>")
            style = f' style="{";".join(st)}"' if st else ""
            tds.append(f"<td{attrs}{style}>{text}</td>")
        h_style = f' style="height:{round(ht * 4 / 3)}px"' if ht else ""
        rows.append(f'<tr{h_style}><th class="rh">{r}</th>{"".join(tds)}</tr>')
    head = "".join(f'<th data-c="{get_column_letter(c)}">{get_column_letter(c)}</th>' for c in range(1, max_c + 1))
    return (f'<table style="width:{total}px"><colgroup>{"".join(cols)}</colgroup>'
            f'<thead><tr><th class="corner"></th>{head}</tr></thead><tbody>{"".join(rows)}</tbody></table>')


def plural(n):
    return f"{n} sheet{'s' if n != 1 else ''}"


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


PAGE_CSS = """
:root{color-scheme:light dark;--brand:#16294a;--brass:#c9a24a;--chrome:#f1f2f5;--hdrline:#d3d7de;--ink:#1d2433;--muted:#5b6678;--hl:#dde5f2;--hover:#e4e7ed;--tabon:#fff;--page:#fff}
@media (prefers-color-scheme:dark){:root{--brand:#0e1a30;--chrome:#18233a;--hdrline:#2c3a56;--ink:#e8ecf3;--muted:#9aa7bf;--hl:#2a3b5e;--hover:#223049;--tabon:#0b1424;--page:#0b1424}}
*{box-sizing:border-box}
html{-webkit-text-size-adjust:100%;-webkit-tap-highlight-color:transparent;overscroll-behavior:none}
html,body{height:100%}
body{margin:0;height:100dvh;display:flex;flex-direction:column;background:var(--page);color:var(--ink);font-family:"Segoe UI",system-ui,-apple-system,Arial,sans-serif;font-size:13px}
.title{display:flex;align-items:center;gap:6px;padding:max(8px,env(safe-area-inset-top)) max(12px,env(safe-area-inset-right)) 8px max(12px,env(safe-area-inset-left));background:var(--brand);color:#fff;border-bottom:2px solid var(--brass)}
.title a{color:#fff;text-decoration:none;opacity:.85}
.title .back{display:grid;place-items:center;width:40px;height:40px;margin-left:-8px;font-size:18px;border-radius:6px}
.title h1{font-size:15px;font-weight:600;margin:0;flex:1;min-width:0;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
.title h1 span{font-weight:400;opacity:.75}
.title .dl{display:inline-flex;align-items:center;min-height:36px;border:1px solid rgba(255,255,255,.55);border-radius:6px;padding:0 12px;white-space:nowrap;font-weight:600;opacity:1;transition-property:scale,background-color;transition-duration:150ms;transition-timing-function:ease-out}
.title .dl:active,.title .back:active{scale:.96;background:rgba(255,255,255,.14)}
.fbar{display:flex;border-bottom:1px solid var(--hdrline);height:30px;background:var(--page)}
.fbar .name{width:90px;border-right:1px solid var(--hdrline);padding:6px 8px;font-family:Consolas,ui-monospace,monospace}
.fbar .fx{padding:6px 10px;color:var(--muted);font-style:italic;border-right:1px solid var(--hdrline)}
.fbar .val{flex:1;padding:6px 10px;font-family:Consolas,ui-monospace,monospace;white-space:nowrap;overflow-x:auto;scrollbar-width:none}
.grid{flex:1;overflow:auto;overscroll-behavior:contain;background:#fff}
.sheet{display:none}
.sheet.on{display:block}
table{border-collapse:separate;border-spacing:0;table-layout:fixed;font-family:Arial,Helvetica,sans-serif;font-size:10pt;color:#000}
tr{height:21px}
td{padding:2px 4px;white-space:nowrap;vertical-align:bottom;overflow:visible;cursor:cell}
td.sel{outline:2px solid var(--brand);outline-offset:-1px}
th{background:#f1f2f5;color:#555;font:400 11px "Segoe UI",Arial,sans-serif;box-shadow:inset -1px -1px 0 #d3d7de;padding:0;user-select:none}
thead th{position:sticky;top:0;z-index:2;height:22px}
th.rh{position:sticky;left:0;z-index:1;text-align:center}
th.corner{left:0;z-index:3}
th.on{background:#dde5f2;color:#16294a;font-weight:600;box-shadow:inset -1px -1px 0 #d3d7de,inset 0 -2px 0 #16294a}
th.rh.on{box-shadow:inset -1px -1px 0 #d3d7de,inset -2px 0 0 #16294a}
.tabs{display:flex;background:var(--chrome);border-top:1px solid var(--hdrline);min-height:36px;padding-bottom:env(safe-area-inset-bottom)}
.tabs nav{display:flex;overflow-x:auto;flex:1;scrollbar-width:none}
.tabs button{border:0;border-right:1px solid var(--hdrline);background:transparent;padding:0 16px;font:inherit;color:var(--muted);cursor:pointer;white-space:nowrap;-webkit-user-select:none;user-select:none;transition-property:background-color,color;transition-duration:150ms;transition-timing-function:ease}
.tabs button:active{background:var(--hover)}
.tabs button.on{background:var(--tabon);color:var(--ink);font-weight:600;box-shadow:inset 0 -3px 0 var(--brass)}
.tabs small{align-self:center;color:var(--muted);padding:0 14px;white-space:nowrap}
.zoom{display:flex;align-items:center;gap:2px;padding-inline:8px;border-inline-start:1px solid var(--hdrline)}
.zoom button{border:0;background:transparent;font:inherit;color:var(--ink);min-width:30px;height:32px;padding-inline:6px;border-radius:4px;cursor:pointer;-webkit-user-select:none;user-select:none}
.zoom button:active{background:var(--hover)}
.zoom .pct{min-width:52px;font-variant-numeric:tabular-nums}
.zoom .fit{color:var(--brand);font-weight:600}
@media (prefers-color-scheme:dark){.zoom .fit{color:var(--brass)}}
button:focus-visible,.title a:focus-visible{outline:2px solid var(--brass);outline-offset:-2px}
@media (hover:hover) and (pointer:fine){.tabs button:hover,.zoom button:hover{background:var(--hover)}.title .dl:hover,.title .back:hover{background:rgba(255,255,255,.12)}}
@media (pointer:coarse){.tabs button{min-height:44px}.zoom button{min-width:44px;height:44px}.fbar{height:34px}}
@media (max-width:700px){
  .tabs{flex-wrap:wrap}
  .tabs small,.title h1 span{display:none}
  .tabs nav{flex:1 1 100%}
  .zoom{flex:1 1 100%;justify-content:space-between;border-inline-start:0;border-top:1px solid var(--hdrline);padding-inline:max(8px,env(safe-area-inset-left)) max(8px,env(safe-area-inset-right))}
  .fbar .name{width:64px}
}
@media (prefers-reduced-motion:reduce){.title .dl:active,.title .back:active{scale:1}}
"""

PAGE_JS = """
const sheets=[...document.querySelectorAll('.sheet')],tabs=[...document.querySelectorAll('.tabs nav button')];
const grid=document.querySelector('.grid'),nameBox=document.querySelector('.fbar .name'),valBox=document.querySelector('.fbar .val');
let sel=null,hl=[];
function select(td){
  if(sel)sel.classList.remove('sel');
  hl.forEach(h=>h&&h.classList.remove('on'));
  sel=td;hl=[];
  if(!td){nameBox.textContent='';valBox.textContent='';return}
  td.classList.add('sel');
  const a=td.dataset.a,col=a.match(/[A-Z]+/)[0];
  nameBox.textContent=a;
  valBox.textContent=td.dataset.f??td.textContent;
  hl=[td.closest('.sheet').querySelector('th[data-c="'+col+'"]'),td.parentElement.firstElementChild];
  hl.forEach(h=>h&&h.classList.add('on'));
}
const ZMIN=0.25,ZMAX=2,STEP=0.1,pctBtn=document.querySelector('.zoom .pct');
let zoom=1;
function table(){return document.querySelector('.sheet.on table')}
function setZoom(z){
  zoom=Math.min(ZMAX,Math.max(ZMIN,Math.round(z*100)/100));
  table().style.zoom=zoom;
  const pct=Math.round(zoom*100)+'%';
  pctBtn.textContent=pct;pctBtn.setAttribute('aria-label','Zoom '+pct+', reset to 100%');
}
function fit(){
  const t=table();t.style.zoom=1;
  const head=t.querySelector('thead').offsetHeight;
  setZoom(Math.min(grid.clientWidth/t.offsetWidth,(grid.clientHeight-head)/(t.offsetHeight-head),1));
}
const narrow=matchMedia('(max-width:700px)');
function fitWidth(){const t=table();t.style.zoom=1;setZoom(Math.max(.5,Math.min(1,grid.clientWidth/t.offsetWidth)))}
document.querySelector('.zoom .out').onclick=()=>setZoom(zoom-STEP);
document.querySelector('.zoom .in').onclick=()=>setZoom(zoom+STEP);
pctBtn.onclick=()=>setZoom(1);
document.querySelector('.zoom .fit').onclick=fit;
function show(i){
  sheets.forEach((s,j)=>s.classList.toggle('on',i===j));
  tabs.forEach((t,j)=>t.classList.toggle('on',i===j));
  grid.scrollTo(0,0);select(null);narrow.matches?fitWidth():setZoom(zoom);
  history.replaceState(null,'','#'+sheets[i].id);
}
tabs.forEach((t,i)=>t.onclick=()=>show(i));
grid.addEventListener('click',e=>{const td=e.target.closest('td');if(td)select(td)});
const start=sheets.findIndex(s=>'#'+s.id===location.hash);
show(start<0?0:start);
"""


def build_workbook(name):
    path = DOCS / f"{name}.xlsx"
    wb = openpyxl.load_workbook(path, data_only=True, rich_text=True)
    fwb = openpyxl.load_workbook(path, rich_text=True)  # same file, formulas instead of cached values
    theme = theme_colors(wb)
    sheets = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
    tabs = "".join(f'<button type="button">{html.escape(ws.title)}</button>' for ws in sheets)
    body = "".join(f'<div class="sheet" id="{slug(ws.title)}">{render_sheet(ws, theme, fwb[ws.title])}</div>' for ws in sheets)
    xlsx = f"{slug(name)}.xlsx"
    shutil.copy(path, SITE / xlsx)
    page = f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">
<meta name="theme-color" content="#16294a">
<title>{html.escape(name)} | {OWNER}</title><style>{PAGE_CSS}</style></head><body>
<div class="title"><a class="back" href="index.html" aria-label="Back to portfolio">&larr;</a><h1>{html.escape(name)}.xlsx <span>&middot; {OWNER}</span></h1>
<a class="dl" href="{xlsx}" download>Download</a></div>
<div class="fbar"><div class="name"></div><div class="fx">fx</div><div class="val"></div></div>
<div class="grid">{body}</div>
<div class="tabs"><nav>{tabs}</nav><small>Portfolio piece &middot; not investment advice</small><div class="zoom" role="group" aria-label="Zoom"><button type="button" class="out" aria-label="Zoom out">&minus;</button><button type="button" class="pct" aria-label="Zoom 100%, reset to 100%">100%</button><button type="button" class="in" aria-label="Zoom in">+</button><button type="button" class="fit" title="Fit the whole sheet on screen">Fit</button></div></div>
<script>{PAGE_JS}</script></body></html>"""
    (SITE / f"{slug(name)}.html").write_text(page, encoding="utf-8")
    return sheets[0].title, len(sheets)


# Portfolio home: profile copy lives in index.template.html; the model register is filled from the workbooks.
# Each output is (cell on the file's first sheet, format or None for text cells, caption).
CASES = [
    ("Monde Nissin Corporation valuation",
     "What is Monde Nissin worth? Three methods built on the same company financials. Share price at the time: {price}.",
     [("Comparable Company Analysis", "Benchmarks Monde against Universal Robina, Century Pacific, and San Miguel Food and Beverage",
       ("D23", None, "LTM implied enterprise value, ₱ in millions")),
      ("Precedent Transaction Analysis", "Values Monde on prices paid in Philippine and ASEAN food and beverage acquisitions",
       ("D19", None, "LTM implied enterprise value, ₱ in millions")),
      ("Discounted Cash Flow Analysis", "Values Monde on projected free cash flow through 2031",
       ("P14", None, "sensitivity-implied enterprise value (exit multiple method)"))]),
    ("Megaworld acquisition of Filinvest Land",
     "Would buying Filinvest Land raise or lower Megaworld's earnings per share?",
     [("Merger Consequences Analysis", "Tests a deal funded with cash, new debt, and Megaworld shares",
       ("U15", "{:+.1%}", "EPS accretion, FY 2026E"))]),
]


def cell_text(name, ref, fmt):
    v = openpyxl.load_workbook(DOCS / f"{name}.xlsx", data_only=True).worksheets[0][ref].value
    if fmt:
        return fmt.format(v)
    return re.sub(r"(\d{1,3}(?:,\d{3})+)\.\d+", r"\1", str(v))  # drop centavos on millions


RANGE = re.compile(r"₱([\d,]+) - ₱([\d,]+)")


def football_field(models, outs):
    """Football-field chart: one bar per method on a shared scale. Empty unless every output is a ₱ range."""
    spans = [RANGE.fullmatch(o) for o in outs]
    if len(models) < 2 or not all(spans):
        return ""
    nums = [tuple(float(g.replace(",", "")) for g in m.groups()) for m in spans]
    step = 10_000
    lo = int(min(a for a, _ in nums) * 0.92 // step * step)
    hi = int(-(-max(b for _, b in nums) * 1.05 // step) * step)
    rows = "".join(
        f'<div class="f-row"><span class="f-name">{html.escape(n.removesuffix(" Analysis"))}</span>'
        f'<span class="f-track" aria-hidden="true"><i style="left:{(a - lo) / (hi - lo):.2%};width:{(b - a) / (hi - lo):.2%}"></i></span>'
        f'<span class="f-val">{o.replace(" - ", " &ndash; ")}</span></div>'
        for (n, *_), (a, b), o in zip(models, nums, outs))
    ticks = "".join(f"<span>₱{v:,.0f}</span>" for v in (lo, (lo + hi) / 2, hi))
    return (f'<figure class="field"><figcaption>Implied enterprise value range, ₱ in millions</figcaption>{rows}'
            f'<div class="f-axis" aria-hidden="true"><div class="f-ticks">{ticks}</div></div></figure>')


def build_index():
    price = cell_text("Comparable Company Analysis", "C7", "₱{:.2f}")
    blocks = []
    for title, intro, models in CASES:
        outs = [cell_text(n, ref, fmt) for n, _, (ref, fmt, _) in models]
        field = football_field(models, outs)
        rows = "".join(
            f'<a class="model" href="{slug(n)}.html"><span class="main"><b>{html.escape(n)}</b><span class="ask">{html.escape(ask)}</span></span>'
            f'<span class="out{" gain" if o[0] == "+" else " loss" if o[0] == "-" else ""}">{html.escape(o.replace(" - ", " – "))}<small>{html.escape(cap)}</small></span>'
            f'<span class="open">Open model</span></a>'
            for (n, ask, (_, _, cap)), o in zip(models, outs))
        blocks.append(f'  <div class="case"><div class="case-head"><h3>{html.escape(title)}</h3><p>{html.escape(intro.format(price=price))}</p></div>{field}{rows}</div>')
    count = sum(len(models) for _, _, models in CASES)
    page = ((ROOT / "index.template.html").read_text(encoding="utf-8")
            .replace("{{MODELS}}", "\n".join(blocks)).replace("{{MODEL_COUNT}}", str(count)))
    (SITE / "index.html").write_text(page, encoding="utf-8")
    shutil.copy(ROOT / "assets" / "jeuz-salvedia-icon.jfif", SITE / "jeuz-salvedia.jpg")
    shutil.copy(DOCS / "Jeuz_Salvedia_Resume.pdf", SITE / "Jeuz_Salvedia_Resume.pdf")


if __name__ == "__main__":
    selfcheck()
    SITE.mkdir(exist_ok=True)
    for n in FILES:
        first, count = build_workbook(n)
        print(f"built {slug(n)}.html ({plural(count)})")
    build_index()
    print("built index.html")
